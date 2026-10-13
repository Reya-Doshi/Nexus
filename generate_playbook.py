import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Colors
C_NAVY = RGBColor(16, 44, 87)
C_BLUE = RGBColor(53, 110, 186)
C_DARK = RGBColor(33, 37, 41)
C_GOLD = RGBColor(180, 83, 9)

def add_header(title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(title + '\n')
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = C_NAVY
    
    r2 = p.add_run(subtitle)
    r2.font.size = Pt(12)
    r2.font.italic = True
    r2.font.color.rgb = C_BLUE

def add_heading1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = C_NAVY

def add_heading2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = C_BLUE

def add_bullet(bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(bold_prefix + ': ')
    r1.font.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)

def add_qa(q_num, question, answer_bullets):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(8)
    p_q.paragraph_format.space_after = Pt(2)
    r_q = p_q.add_run(f'Q{q_num}: {question}')
    r_q.font.bold = True
    r_q.font.size = Pt(11)
    r_q.font.color.rgb = C_NAVY
    
    for ans in answer_bullets:
        p_a = doc.add_paragraph(style='List Bullet')
        p_a.paragraph_format.space_after = Pt(1)
        r_a = p_a.add_run(ans)
        r_a.font.size = Pt(10)

def add_box(title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(f'[{title}] ')
    r1.font.bold = True
    r1.font.color.rgb = C_GOLD
    r2 = p.add_run(text)
    r2.font.italic = True
    r2.font.size = Pt(10.5)

# --- COVER / HEADER ---
add_header('INFOSYS IGNITE ENG NEXT 2026: WINNING PLAYBOOK', 'Problem Statement NC-04: Self-Healing Network Simulator (Project ReRoute)\nQualcomm Cloud AI Playground & Snapdragon NPU Acceleration')

add_box('TARGET PROBLEM STATEMENT NC-04 SUMMARY', 'NC-04 requires an AI-assisted tool that continuously monitors telemetry, catches degradation/failures, pinpoints root cause, automatically reroutes traffic, and visually recovers the network in real time. Platforms: Qualcomm Cloud AI Playground (Cloud AI 100) + Snapdragon X Elite NPU.')

# ==================== SECTION 1: MASTER 3-MINUTE PITCH ====================
add_heading1('1. THE EXACT 3-MINUTE PITCH SCRIPT (CALIBRATED FOR NC-04)')

add_heading2('A. 0:00 - 0:45 | PRANATHI (The Problem & The Closed-Loop Challenge)')
add_bullet('Exact Pitch Script', '"Good morning, judges. In modern enterprise and carrier networks, unanticipated link degradation, packet drops, and node failures trigger catastrophic cascading downtime before human operators can intervene. Traditional network operations are purely reactive: NOC teams receive alerts only after end-users report degradation, requiring manual log parsing that can take minutes to hours. Testing automated remediation on live production backbones is prohibitively risky. What telecom and data center networks desperately need is an autonomous Detect-Diagnose-Heal closed loop—validated safely in a simulated environment before production deployment. That is the exact mission of our project for NC-04: ReRoute—an AI-driven Self-Healing Network Simulator."')
add_bullet('Delivery Strategy', 'Clear, authoritative delivery. Hand over at 0:42 with: "Swathi will now walk you through our Qualcomm-powered self-healing architecture."')

add_heading2('B. 0:45 - 1:30 | SWATHI (The Architecture & Qualcomm Cloud AI Integration)')
add_bullet('Exact Pitch Script', '"To solve NC-04, we engineered an autonomous self-healing telemetry pipeline powered by Qualcomm AI silicon. Our system is built on three synchronized pillars: First, continuous real-time telemetry streaming monitoring latency, jitter, and link saturation across network topology nodes. Second, a cognitive root-cause diagnosis engine running Meta\'s Llama-3.1-8B directly on Qualcomm Cloud AI 100 via the Cirrascale Qualcomm AI Playground API. The moment an anomaly is detected, our model identifies the root cause—such as optical link degradation or buffer bloat—and formulates the optimal reroute policy. Third, an automated corrective execution layer using FastAPI and dynamic graph optimization that reroutes around the degraded link and fires instant Telegram incident advisories to network engineers. Reya will now demonstrate the live fault-injection and self-healing recovery."')
add_bullet('Delivery Strategy', 'Emphasize "Detect-Diagnose-Heal". Hand over at 1:28 with: "Reya will now demonstrate the live fault-injection and self-healing recovery."')

add_heading2('C. 1:30 - 2:55 | REYA (Live Prototype Demo + Snapdragon NPU Benchmark Flex)')
add_bullet('Exact Pitch Script', '"Thank you! Here is our live Self-Healing Network Dashboard. (Trigger fault injection on web UI). As you see, we just injected a link failure into Node 3. Immediately, our telemetry watcher catches the anomaly. In fractions of a second, the backend queries our Qualcomm Cloud AI 100 endpoint—and here is the live AI diagnosis: link degradation identified, mitigating by re-weighting routing topology. The traffic visually reroutes around the damaged corridor, and our Telegram notification bot simultaneously alerts the NOC team with root-cause analysis—zero packet loss, zero user disruption. (Switch to Tab 2: Qualcomm Workbench). Furthermore, in real-world switches and edge SDN routers, running telemetry anomaly detection cannot suffer cloud round-trip delay. We deployed our telemetry perception model to the Qualcomm AI Hub, compiled it with QNN, and executed it on a physical Snapdragon X Elite Compute Device. As confirmed on our live Qualcomm Workbench dashboard—job jglw1r28p—our model runs in just 0.64 milliseconds on the Snapdragon NPU. That is line-rate anomaly detection at over 1,500 inferences per second on physical Qualcomm silicon. Thank you!"')
add_bullet('Delivery Strategy', 'Show the visual reroute on Tab 1, then snap to Tab 2 (Qualcomm Workbench). Finish cleanly at 2:55.')

doc.add_page_break()

# ==================== SECTION 2: PRANATHI ====================
add_heading1('2. PRANATHI: NC-04 TELEMETRY & NETWORK INTELLIGENCE')
add_heading2('Role: Telemetry Dynamics & Network Anomaly Specialist')
add_bullet('Key Focus', 'Telemetry signals (latency, jitter, packet loss, throughput), synthetic fault injection mechanisms, and early anomaly detection thresholds.')

add_heading2('Pranathi\'s 10000x Questionnaire Defense')

add_qa(1, 'How does your simulator inject and detect faults before end users feel impact?', [
    'We simulate network telemetry metrics using continuous sliding windows monitoring round-trip time (RTT), jitter variance, and synthetic probe drops.',
    'Before a total link failure occurs, telemetry displays micro-anomalies (gradual RTT jitter rise and TCP retransmission bursts). Our model flags these pre-failure degradation signatures before active packet drops hit customer payloads.'
])

add_qa(2, 'What network topologies does the simulator support?', [
    'The simulator is topology-agnostic. In our demo, we model an enterprise multi-mesh topology with redundant paths, spine-leaf core interconnects, and edge ingress routers.'
])

add_qa(3, 'How do you avoid false positive self-healing triggers during normal traffic spikes?', [
    'We distinguish transient burstiness from persistent degradation by tracking multi-metric cross-correlation: A normal video stream spike increases throughput with stable low jitter, whereas actual hardware or fiber degradation shows rising jitter and packet drops regardless of throughput.'
])

doc.add_page_break()

# ==================== SECTION 3: SWATHI ====================
add_heading1('3. SWATHI: QUALCOMM CLOUD AI 100 & AUTOMATED REMEDIATION')
add_heading2('Role: System Architect, Qualcomm Cloud AI 100 & Remediation Orchestrator')
add_bullet('Key Focus', 'Qualcomm Cloud AI Playground API integration, Llama-3.1-8B root cause reasoning, dynamic rerouting solver, and Telegram alert webhooks.')

add_heading2('Swathi\'s 10000x Questionnaire Defense')

add_qa(1, 'How exactly did you use the Qualcomm Cloud AI Playground for NC-04?', [
    'NC-04 explicitly specifies the Qualcomm Cloud AI Playground. We configured our FastAPI backend to communicate directly with Qualcomm Cloud AI 100 instances hosted on Cirrascale (v2 API).',
    'We stream formatted telemetry logs into Meta\'s Llama-3.1-8B on Qualcomm Cloud AI 100 to generate deterministic root-cause diagnosis and actionable JSON mitigation parameters.'
])

add_qa(2, 'How does the network automatically "reroute around the fault" in real time?', [
    'We implement dynamic graph re-weighting (modified Dijkstra / Shortest Path First).',
    'The moment Qualcomm AI confirms degradation on Edge (U, V), our engine dynamically assigns infinite cost to that edge, instantly triggering routing table convergence to shift active traffic flows through healthy alternate paths.'
])

add_qa(3, 'Why include Telegram bot alerts in a self-healing system?', [
    'While the system heals autonomously, network administrators must maintain auditability and post-incident visibility. Telegram delivers instant pushed forensic diagnostics (time of failure, affected node, root cause, and remediation applied) directly to engineers on call.'
])

doc.add_page_break()

# ==================== SECTION 4: REYA ====================
add_heading1('4. REYA: QUALCOMM SILICON LEAD & HARDWARE BENCHMARKS')
add_heading2('Role: Lead Architect, Hardware Acceleration & Qualcomm NPU Specialist')
add_bullet('Key Focus', 'Snapdragon X Elite NPU execution, Qualcomm AI Hub SDK, QNN DLC compilation, line-rate hardware latency, and end-to-end technical leadership.')

add_heading2('Reya\'s 10000x Questionnaire Defense')

add_qa(1, 'The problem statement mentions Qualcomm Cloud AI Playground. Why did you also test on Snapdragon X Elite?', [
    'Qualcomm Cloud AI Playground is ideal for centralized orchestration and large LLM root-cause reasoning. However, in next-gen telecom (5G UPF, SDN switches, autonomous routers), waiting for cloud round-trip latency to detect a line failure takes 100-300 ms—during which millions of packets are lost.',
    'By compiling our anomaly detector for the Snapdragon X Elite Hexagon NPU, we proved that network equipment can detect and isolate failures on-box in just 0.64 ms, creating the ultimate hybrid Edge-to-Cloud architecture.'
])

add_qa(2, 'Show us proof that this actually ran on Qualcomm hardware.', [
    'We submitted and verified our job through Qualcomm AI Hub on a physical Snapdragon X Elite Compute Reference Device (CRD).',
    'Our official job ID is jglw1r28p, verifiable live on Qualcomm Workbench at: https://workbench.aihub.qualcomm.com/jobs/jglw1r28p/',
    'The report confirms zero operator fallback and a hardware inference latency of 0.64 milliseconds (641 microseconds).'
])

add_qa(3, 'How does 0.64 ms inference benefit carrier-grade network throughput?', [
    'At 0.64 ms, our model can execute over 1,560 health-check inferences per second per NPU core.',
    'This allows real-time packet-window telemetry inspection at line speed without creating CPU bottlenecks or degrading router throughput.'
])

add_qa(4, 'What was your compilation process using Qualcomm AI Hub SDK?', [
    'We serialized our model using PyTorch tracing with static input dimensions matching our telemetry feature matrix.',
    'We submitted a compilation job to Qualcomm AI Hub with --target_runtime qnn_dlc, fusing operations for the Hexagon Tensor Processor (HTP).',
    'We then executed a profile job on physical Snapdragon hardware to extract real-world execution latency, memory footprint, and NPU execution percentage.'
])

out_path = 'PROJECT_REROUTE_WINNING_PLAYBOOK.docx'
doc.save(out_path)
print('SUCCESS! Document updated at:', out_path)
